# ======================================================================
# IMPORTS
# ======================================================================

import re
import fitz  # PyMuPDF
import json
import datetime
from io import BytesIO
from typing import Dict, Any, Optional

from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table

from flask import render_template
from weasyprint import HTML
import openpyxl


# ======================================================================
# EXTRATOR PADRÃO (VERSÃO A – SIMPLES, SEGURO E TOLERANTE)
# ======================================================================

def extrair_padrao_pdf(texto: str):
    flags = re.IGNORECASE | re.DOTALL

    dados = {
        "Cliente": "N/A",
        "CPF": "N/A",
        "Telefone": "N/A",
        "Data_de_Pagamento": "N/A",
        "FormaDePagamento": "N/A",
        "Data_do_Evento": "N/A",
        "Local_do_Evento": "N/A",
        "Valor_Total_do_Pedido": "N/A",
        "produtosContratadosJson": "[]",
        "Responsavel": "N/A",
        "Como nos conheceu": "N/A"
    }

    # ==============================
    # CLIENTE
    # ==============================
    m = re.search(r"CONTRATANTE:\s*(?:Sr\(a\))?\s*(.*?)[,\n]", texto, flags)
    if m:
        dados["Cliente"] = m.group(1).strip()

    # ==============================
    # CPF
    # ==============================
    m = re.search(r"CPF[:\s]*([\d.\-]+)", texto)
    if m:
        dados["CPF"] = m.group(1).strip()

    # ==============================
    # TELEFONE
    # ==============================
    m = re.search(
        r"Tel\.?\s*\(?\d{2}\)?\s*\d[\d.\- ]+\d",
        texto,
        flags
    )
    if m:
        telefone = m.group(0).strip().rstrip(".")
        dados["Telefone"] = telefone




    # ==============================
    # VALOR TOTAL REAL
    # ==============================
    m = re.search(
    r"valor total de R\$\s*([\d.,]+)",
    texto, re.IGNORECASE
)
    if m:
        dados["Valor_Total_do_Pedido"] = m.group(1)


    # =================================================================
    # PRODUTOS — SUPORTE COMPLETO PARA TABELA QUEBRADA DO SEU PDF
    # =================================================================
    produtos = []

    pattern = r"(\d{1,4})\s*\n([A-Za-zÀ-ÿ0-9 \-\.]+)\s*\n([\d.,]+)\s*\n([\d.,]+)"
    matches = re.findall(pattern, texto)

    for qtd, nome, vunit, vtotal in matches:
        produtos.append({
        "Quantidade": qtd,
        "Produto": nome.strip(),
        "Valor Unitário": vunit,
        "Valor Total Item": vtotal
    })

    dados["produtosContratadosJson"] = json.dumps(produtos, ensure_ascii=False)


    # ==============================
    # PAGAMENTO (data + forma)
    # ==============================
    m = re.search(
    r"pagos?\s+no\s+dia\s*\n?\s*([\d/\-]{8,10})\s*(a vista|à vista|pix|cart[aã]o|dinheiro|dep[oó]sito|transfer[eê]ncia)?",
    texto,
    re.IGNORECASE
)
    if m:
        dados["Data_de_Pagamento"] = m.group(1).strip()
        forma = m.group(2)
        dados["FormaDePagamento"] = forma.strip() if forma else "N/A"


    # ==============================
    # DATA DO EVENTO
    # ==============================
    m = re.search(r"O evento acontecerá no dia[:\s]*([\d/\-]{8,10})", texto, flags)
    if m:
        dados["Data_do_Evento"] = m.group(1).strip()

    # ==============================
    # LOCAL DO EVENTO
    # ==============================
    m = re.search(
    r"Local do evento:\s*(.*)",
    texto, re.IGNORECASE
)
    if m:
        dados["Local_do_Evento"] = m.group(1).strip(" –-")


    # ==============================
    # COMO NOS CONHECEU
    # ==============================
    m = re.search(r"Como nos conheceu[:\s]*(.*)", texto)
    if m:
        dados["Como nos conheceu"] = m.group(1).strip()

    # ==============================
    # RESPONSÁVEL
    # ==============================
    m = re.search(r"RESPONSÁVEL PELO CONTRATO[:\s]*(.*)", texto)
    if m:
        dados["Responsavel"] = m.group(1).strip()

        print("DEBUG FINAL:", json.dumps(dados, ensure_ascii=False, indent=2))

    return dados

    



def extrair_padrao_docx(texto: str):
    flags = re.IGNORECASE | re.DOTALL

    dados = {
        "Cliente": "N/A",
        "CPF": "N/A",
        "Telefone": "N/A",
        "Data_de_Pagamento": "N/A",
        "FormaDePagamento": "N/A",
        "Data_do_Evento": "N/A",
        "Local_do_Evento": "N/A",
        "Valor_Total_do_Pedido": "N/A",
        "produtosContratadosJson": "[]",
        "Responsavel": "N/A",
        "Como nos conheceu": "N/A"
    }

    # Normalizar texto para remover múltiplos espaços e quebras
    texto_norm = re.sub(r"[ \t]+", " ", texto)
    texto_norm = re.sub(r"\n+", "\n", texto_norm)

    # =========================================================
    # CLIENTE
    # =========================================================
    m = re.search(r"CONTRATANTE:\s*(?:Sr\(a\))?\s*(.*?)[,\n]", texto_norm, flags)
    if m:
        dados["Cliente"] = m.group(1).strip()

    # CPF
    m = re.search(r"CPF[:\s]*([\d.\-]+)", texto_norm)
    if m:
        dados["CPF"] = m.group(1).strip()

    # TELEFONE
    m = re.search(r"Tel\.?\s*\(?\d{2}\)?\s*\d[\d.\- ]+\d", texto_norm, flags)
    if m:
        dados["Telefone"] = m.group(0).strip().rstrip(".")

    # =========================================================
    # DATA DO EVENTO
    # =========================================================
    m = re.search(r"O evento acontecerá no dia[:\s]*([\d/\-]{8,10})", texto_norm)
    if m:
        dados["Data_do_Evento"] = m.group(1)

    # LOCAL DO EVENTO
    m = re.search(r"Local do evento[:\s]*([^\n]+)", texto_norm)
    if m:
        dados["Local_do_Evento"] = m.group(1).strip(" –-")

    # =========================================================
    # VALOR TOTAL
    # =========================================================
    m = re.search(r"valor total de R\$\s*([\d.,]+)", texto_norm, flags)
    if m:
        dados["Valor_Total_do_Pedido"] = m.group(1)

    # =========================================================
    # DATA + FORMA DE PAGAMENTO (Versão Universal)
    # =========================================================

    # 1) Capturar TODAS as datas de pagamento
    datas_pagamento = re.findall(
        r"(?:pago|pagos|pagamento|quitado|quitada|pagou)[^0-9]{0,20}(\d{2}/\d{2}/\d{4})",
        texto_norm,
        flags
)

    # Se houver mais de uma, use a última (data de quitação final)
    if datas_pagamento:
        dados["Data_de_Pagamento"] = datas_pagamento[-1]
    else:
        dados["Data_de_Pagamento"] = "N/A"


    # 2) Capturar forma de pagamento em QUALQUER PARTE DO TEXTO
    m_forma = re.search(
        r"(pix|cart[aã]o(?: de cr[eé]dito| de d[eé]bito)?|dinheiro|dep[oó]sito|transfer[eê]ncia|boleto)",
    texto_norm,
    flags
)

    if m_forma:
        dados["FormaDePagamento"] = m_forma.group(1).strip()
    else:
        dados["FormaDePagamento"] = "N/A"


    # =========================================================
    # PRODUTOS (DOCX — LINHAS SEMPRE EM COLUNA)
    # =========================================================
    produtos = []

    linhas = texto_norm.split("\n")
    buffer = []

    for ln in linhas:
        if re.fullmatch(r"\d+", ln.strip()):
            # Começou um produto
            buffer = [ln.strip()]
        elif buffer and len(buffer) == 1:
            buffer.append(ln.strip())
        elif buffer and len(buffer) == 2:
            buffer.append(ln.strip())
        elif buffer and len(buffer) == 3:
            buffer.append(ln.strip())
            # agora temos 4 linhas = produto completo
            produtos.append({
                "Quantidade": buffer[0],
                "Produto": buffer[1],
                "Valor Unitário": buffer[2],
                "Valor Total Item": buffer[3]
            })
            buffer = []

    dados["produtosContratadosJson"] = json.dumps(produtos, ensure_ascii=False)

    # =========================================================
    # Responsável / Como conheceu
    # =========================================================
    m = re.search(r"RESPONSÁVEL PELO CONTRATO[:\s]*(.*)", texto_norm)
    if m:
        dados["Responsavel"] = m.group(1).strip()

    m = re.search(r"Como nos conheceu[:\s]*(.*)", texto_norm)
    if m:
        dados["Como nos conheceu"] = m.group(1).strip()

    print("DEBUG DOCX FINAL:", json.dumps(dados, ensure_ascii=False, indent=2))

    return dados



# ======================================================================
# EXTRATOR DO SISTEMA (REGEX OFICIAL – VERSÃO ESTÁVEL)
# ======================================================================

def _extrair_com_regex(texto: str) -> Dict[str, Any]:
    """
    Extrator completo do CONTRATO DO SISTEMA.
    Mantém compatibilidade total com o contrato gerado pelo DOCX oficial.
    """

    print("[INFO] Executando extração REGEX (Sistema)...")
    flags = re.DOTALL | re.IGNORECASE

    # Estrutura padrão de retorno
    dados = {
        "Contratante": {
            "Nome": "N/A",
            "CPF": "N/A",
            "Telefone": "N/A",
            "Email": "N/A",
            "RG": "N/A",
            "Endereco": "N/A",
        },
        "Data_do_Evento": "N/A",
        "Local_do_Evento": "N/A",
        "produtosContratadosJson": "[]",
        "Data_de_Pagamento": "N/A",
        "Valor_Total_do_Pedido": "N/A",
        "FormaDePagamento": "N/A",
        "Responsavel": "N/A",
        "Como nos conheceu": "N/A"
    }

    # Normalização simples
    texto = texto.replace("\r", "")
    linhas = texto.splitlines()

    # ==========================================================
    # 1) CONTRATANTE
    # ==========================================================
    bloco = re.search(r"CONTRATANTE:\s*Sr\(a\)([\s\S]*?)CONTRATADO", texto, flags)
    if bloco:
        tx = bloco.group(1)

        # Nome
        m = re.search(r"^\s*(.*?)\s*,", tx, flags)
        if m:
            dados["Contratante"]["Nome"] = m.group(1).strip()

        # RG
        m = re.search(r"RG[:\s]*([\d.\- ]+)", tx, flags)
        if m:
            dados["Contratante"]["RG"] = m.group(1).strip()

        # CPF (pontuado ou não)
        m = re.search(r"CPF[:\s]*([\d.\- ]+)", tx, flags)
        if m:
            cpf_raw = re.sub(r"[^\d]", "", m.group(1))
            if len(cpf_raw) == 11:
                dados["Contratante"]["CPF"] = cpf_raw

        # Endereço
        m = re.search(r"domiciliado\(a\)\s*na\s*(.*?)(?:Tel|–|—|$)", tx, flags)
        if m:
            dados["Contratante"]["Endereco"] = m.group(1).strip()

        # Telefone
        m = re.search(r"Tel\.?\s*([\d()\s.\-\/]+)", tx, flags)
        if m:
            dados["Contratante"]["Telefone"] = m.group(1).strip()

        # Email
        m = re.search(r"(?:E-mail|Email)[:\s]*([\w.%+\-]+@[\w.\-]+\.[A-Za-z]{2,})", tx, flags)
        if m:
            dados["Contratante"]["Email"] = m.group(1).strip()

    # Fallback CPF global
    if dados["Contratante"]["CPF"] == "N/A":
        m = re.search(r"\b(\d{11})\b", texto)
        if m:
            dados["Contratante"]["CPF"] = m.group(1)

    # ==========================================================
    # 2) PRODUTOS — Bloco entre CLÁUSULA 1 e CLÁUSULA 2
    # ==========================================================
    print("[INFO] Extraindo produtos (sistema)...")
    produtos = []

    bloco_prod = re.search(r"CLÁUSULA\s*1[\s\S]*?PRODUTOS\s*CONTRATADOS([\s\S]*?)CLÁUSULA\s*2", texto, flags)
    if bloco_prod:
        bruto = bloco_prod.group(1)
        linhas = [l.strip() for l in bruto.splitlines() if l.strip()]

        # Remove possíveis headers
        headers = {"quantidade", "produto", "valor unitário", "valor total", "valor unitario", "total"}
        linhas = [l for l in linhas if l.lower() not in headers]

        i = 0
        while i < len(linhas):
            ln = linhas[i]

            # Ignorar TOTAL
            if ln.lower().startswith("total"):
                i += 1
                continue

            # Caso: quantidade sozinha (linha 14)
            if re.fullmatch(r"\d+", ln):
                qtd = ln
                nome = linhas[i+1] if i+1 < len(linhas) else ""
                vunit = linhas[i+2] if i+2 < len(linhas) else ""
                vtot = linhas[i+3] if i+3 < len(linhas) else ""

                if nome.lower().startswith("total"):
                    i += 1
                    continue

                produtos.append({
                    "Quantidade": qtd,
                    "Produto": nome,
                    "Valor Unitário": vunit,
                    "Valor Total Item": vtot
                })
                i += 4
                continue

            # Caso "14x Brigadeiro"
            m = re.match(r"(\d{1,4})\s*[xX×]\s+(.+)", ln)
            if m:
                produtos.append({
                    "Quantidade": m.group(1),
                    "Produto": m.group(2).strip(),
                    "Valor Unitário": "N/A",
                    "Valor Total Item": "N/A",
                })
                i += 1
                continue

            # Linha não identificada → vira produto genérico
            produtos.append({
                "Quantidade": "N/A",
                "Produto": ln,
                "Valor Unitário": "N/A",
                "Valor Total Item": "N/A",
            })
            i += 1

    dados["produtosContratadosJson"] = json.dumps(produtos, ensure_ascii=False)

    # ==========================================================
    # 3) VALOR TOTAL
    # ==========================================================
    m = re.search(r"TOTAL[:\s]*R?\$?\s*([\d.,]+)", texto, flags)
    if not m:
        m = re.search(r"O valor total de\s*R?\$?\s*([\d.,]+)", texto, flags)
    if m:
        dados["Valor_Total_do_Pedido"] = m.group(1).strip()

    # ==========================================================
    # 4) PAGAMENTO
    # ==========================================================
    m = re.search(r"pagos\s+no\s+dia[:\s]*([\d/]+)\s+([A-Za-z0-9]+)", texto, flags)
    if m:
        dados["Data_de_Pagamento"] = m.group(1)
        dados["FormaDePagamento"] = m.group(2)

    # ==========================================================
    # 5) EVENTO
    # ==========================================================
    m = re.search(r"O evento acontecerá no dia[:\s]*([\d/]+).*?Local do evento[:\s]*(.*)", texto, flags)
    if m:
        dados["Data_do_Evento"] = m.group(1)
        dados["Local_do_Evento"] = m.group(2).split("\n")[0].strip()

    # ==========================================================
    # 6) OUTROS
    # ==========================================================
    m = re.search(r"Como nos conheceu[:\s]*(.*?)\n", texto, flags)
    if m:
        dados["Como nos conheceu"] = m.group(1).strip()

    m = re.search(r"RESPONSÁVEL PELO CONTRATO[:\s]*(.*?)\n", texto, flags)
    if m:
        dados["Responsavel"] = m.group(1).strip()

    print("[INFO] Extração do sistema concluída.")
    return dados


# ======================================================================
# FUNÇÃO DE EXTRAÇÃO DE TEXTO (PDF + DOCX)
# ======================================================================

def _extrair_texto(file_bytes: bytes, file_extension: str) -> Optional[str]:
    """
    Extrai texto de PDF (via PyMuPDF) ou DOCX (via python-docx).
    """

    # Normaliza a extensão
    ext = (file_extension or "").strip().lower()

    # ---------------------------------------
    # PDF
    # ---------------------------------------
    if ext == ".pdf":
        print("[INFO] Extraindo texto de PDF (PyMuPDF)...")
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            texto = ""
            for page in doc:
                texto += page.get_text("text") + "\n"
            return texto
        except Exception as e:
            print(f"[ERRO] Falha ao extrair PDF: {e}")
            return ""

    # ---------------------------------------
    # DOCX
    # ---------------------------------------
    if ext == ".docx":
        print("[INFO] Extraindo texto de DOCX (python-docx)...")
        try:
            texto = ""
            doc_stream = BytesIO(file_bytes)
            document = Document(doc_stream)

            # Cabeçalho
            for section in document.sections:
                for p in section.header.paragraphs:
                    texto += p.text + "\n"

            # Corpo
            body = document.element.body
            for element in body:
                if element.tag.endswith('p'):
                    texto += Paragraph(element, document).text + "\n"
                elif element.tag.endswith('tbl'):
                    table = Table(element, document)
                    for row in table.rows:
                        for cell in row.cells:
                            texto += cell.text + "\n"

            # Rodapé
            for section in document.sections:
                for p in section.footer.paragraphs:
                    texto += p.text + "\n"

            # Normalizações
            texto = texto.replace("–", "-")
            texto = texto.replace("—", "-")
            texto = texto.replace("\t", " ")
            texto = re.sub(r" {2,}", " ", texto)
            texto = re.sub(r"\n{2,}", "\n", texto)

            return texto

        except Exception as e:
            print(f"[ERRO] Falha ao extrair DOCX: {e}")
            return ""

    # ---------------------------------------
    # EXTENSÃO INVÁLIDA
    # ---------------------------------------
    print(f"[ERRO] Extensão não suportada: {ext}")
    return ""







# ======================================================================
# FUNÇÃO CENTRAL: ROTEADOR DE EXTRAÇÃO
# Escolhe automaticamente o EXTRATOR PADRÃO ou SISTEMA
# ======================================================================

def extrair_dados_do_contrato_por_tipo(file_bytes, tipo_analise, file_extension):

    texto = _extrair_texto(file_bytes, file_extension)

    tipo_analise = (tipo_analise or "").lower().strip()

    # ativa extrator do sistema para qualquer variação
    if "sist" in tipo_analise:
        print("[INFO] Usando extrator OFICIAL DO SISTEMA...")
        return extrair_sistema(texto)

    print("[INFO] Usando extrator PADRÃO...")
    ...


    # ----------------------------------------
    # CONTRATO DO SISTEMA
    # ----------------------------------------
    if tipo_analise == "sistema":
        print("[INFO] Usando extrator OFICIAL DO SISTEMA...")
        return extrair_sistema(texto)

    # ----------------------------------------
    # CONTRATO PADRÃO (PDF ou DOCX)
    # ----------------------------------------
    print("[INFO] Usando extrator PADRÃO...")
    if file_extension == ".pdf":
        return extrair_padrao_pdf(texto)
    else:
        return extrair_padrao_docx(texto)









# ======================================================================
# EXTRATOR OFICIAL DO SISTEMA (VERSÃO FINAL ESTÁVEL)
# ======================================================================

def extrair_sistema(texto: str) -> Dict[str, Any]:
    """
    Extrator oficial para contratos GERADOS PELO SISTEMA.
    Extração robusta e precisa, totalmen    te isolada do extrator padrão.
    """

    flags = re.IGNORECASE | re.DOTALL

    dados = {
        "Contratante": {
            "Nome": "N/A",
            "RG": "N/A",
            "CPF": "N/A",
            "Telefone": "N/A",
            "Email": "N/A",
            "Endereco": "N/A"
        },
        "Data_do_Evento": "N/A",
        "Local_do_Evento": "N/A",
        "produtosContratadosJson": "[]",
        "Data_de_Pagamento": "N/A",
        "FormaDePagamento": "N/A",
        "Valor_Total_do_Pedido": "N/A",
        "Responsavel": "N/A",
        "Como nos conheceu": "N/A"
    }

    texto = texto.replace("\r", "")

    # ======================================================================
    # 1) BLOCO DO CONTRATANTE
    # ======================================================================
    bloco = re.search(r"CONTRATANTE[:\s]*(.*?)(?=CONTRATADO)", texto, flags)

    if bloco:
        tx = bloco.group(1)

        # Nome
        m = re.search(r"Sr\(a\)\s*(.*?)[,|\n]", tx)
        if m:
            dados["Contratante"]["Nome"] = m.group(1).strip()

        # RG
        m = re.search(r"RG[:\s]*([\d.\-]+)", tx)
        if m:
            dados["Contratante"]["RG"] = m.group(1).strip()

        # CPF
        m = re.search(r"CPF[:\s]*([\d.\-]+)", tx)
        if m:
            dados["Contratante"]["CPF"] = re.sub(r"\D", "", m.group(1))

        # Email
        m = re.search(r"Email[:\s]*([\w\.-]+@[\w\.-]+\.\w+)", tx)
        if m:
            dados["Contratante"]["Email"] = m.group(1).strip()

        # Telefone
        m = re.search(r"Tel\.?\s*([\d().\-\s]+)", tx)
        if m:
            dados["Contratante"]["Telefone"] = m.group(1).strip()

        # Endereço
        m = re.search(r"domiciliado\(a\)\s*na\s*(.*?)(?=Tel|Email|$)", tx, flags)
        if m:
            dados["Contratante"]["Endereco"] = m.group(1).strip()

    # ======================================================================
    # 2) PRODUTOS – Tabela do sistema
    # ======================================================================
    produtos = []

    bloco_prod = re.search(r"CLÁUSULA\s*1[\s\S]*?PRODUTOS CONTRATADOS([\s\S]*?)CLÁUSULA\s*2", texto, flags)

    if bloco_prod:
        linhas = [
            l.strip()
            for l in bloco_prod.group(1).splitlines()
            if l.strip()
        ]

        # Remover cabeçalhos
        headers = {"quantidade", "produto", "valor unitário", "valor total"}
        linhas = [l for l in linhas if l.lower() not in headers]

        i = 0
        while i < len(linhas):
            ln = linhas[i]

            # Caso 1: tabela vertical (quantidade na linha)
            if re.fullmatch(r"\d+", ln):
                qtd = ln
                nome = linhas[i+1] if i+1 < len(linhas) else "N/A"
                vunit = linhas[i+2] if i+2 < len(linhas) else "N/A"
                vtot = linhas[i+3] if i+3 < len(linhas) else "N/A"

                produtos.append({
                    "Quantidade": qtd,
                    "Produto": nome,
                    "Valor Unitário": vunit,
                    "Valor Total Item": vtot
                })

                i += 4
                continue

            # Caso 2: "14x Brigadeiro"
            m = re.match(r"(\d+)\s*[xX]\s*(.+)", ln)
            if m:
                produtos.append({
                    "Quantidade": m.group(1),
                    "Produto": m.group(2),
                    "Valor Unitário": "N/A",
                    "Valor Total Item": "N/A"
                })
                i += 1
                continue

            i += 1

    dados["produtosContratadosJson"] = json.dumps(produtos, ensure_ascii=False)

    # ================================================================
    # 3) VALOR TOTAL – Captura SOMENTE o "TOTAL: R$ X"
    # ================================================================
    # TOTAL do pedido — capturar apenas linha TOTAL real, não valores quebrados
    m = re.search(r"^TOTAL[:\s]*R?\$?\s*([\d.,]+)$", texto, re.IGNORECASE | re.MULTILINE)
    if m:
        dados["Valor_Total_do_Pedido"] = m.group(1).replace(",", ".")



    

    # ======================================================================
    # 4) PAGAMENTO
    # ======================================================================
        m = re.search(
        r"pagos\s+no\s+dia\s*([\d/\-]{8,10})\s*(PIX|Cartão|Dinheiro|Depósito|Transferência)?",
        texto,
        flags
)
    if m:
        dados["Data_de_Pagamento"] = m.group(1).strip()
        dados["FormaDePagamento"] = m.group(2) or "N/A"


    # ======================================================================
    # 5) DATA & LOCAL DO EVENTO – PARADA AUTOMÁTICA
    # ======================================================================
    m = re.search(
    r"acontecerá no dia[:\s]*([\d/\-]{8,10})\s*-\s*Local do evento[:\s]*(.*?)(?:Como nos conheceu|RESPONSÁVEL|CLÁUSULA|$)",
    texto,
    flags
)
    if m:
        dados["Data_do_Evento"] = m.group(1).strip()
        dados["Local_do_Evento"] = m.group(2).strip()



    # ======================================================================
    # 6) COMO CONHECEU / RESPONSÁVEL
    # ======================================================================
    m = re.search(r"Como nos conheceu[:\s]*(.*)", texto)
    if m:
        dados["Como nos conheceu"] = m.group(1).strip()

    m = re.search(r"RESPONSÁVEL PELO CONTRATO[:\s]*(.*)", texto)
    if m:
        dados["Responsavel"] = m.group(1).strip()

    return dados



    # ==============================================================================
# SEÇÃO 3: GERAÇÃO DE DOCUMENTOS (Inalterado)
# ==============================================================================
# (Todas as funções de gerar_contrato_docx, gerar_contrato_pdf_direto,
# gerar_relatorio_entrega, e exportar_para_excel permanecem 100% iguais)

def gerar_contrato_docx(dados: Dict[str, Any]) -> Optional[BytesIO]:
    try:
        document = Document()
        document.add_heading('Divinos Doces Finos', 0)
        document.add_paragraph('CONTRATO', style='Normal')
        document.add_paragraph()

        # --- CONTRATANTE ---
        contratante = dados.get('Contratante', {})
        contratante_texto = document.add_paragraph()
        contratante_texto.add_run('CONTRATANTE: ').bold = True
        contratante_texto.add_run(
            f"Sr(a) {contratante.get('Nome', 'N/A')}, brasileiro(a), "
            f"portador(a) da cédula de RG: {contratante.get('RG', 'N/A')} e CPF: {contratante.get('CPF', 'N/A')}, "
            f"residente e domiciliado(a) na {contratante.get('Endereco', 'N/A')} - "
            f"Tel. {contratante.get('Telefone', 'N/A')}."
        )

        contratado_texto = document.add_paragraph()
        contratado_texto.add_run('CONTRATADO: ').bold = True
        contratado_texto.add_run(
            "Divinos Doces Finos, inscrito sob o CNPJ: 18.826.801/0001-76, com sede na Rua Curupacê, 392 "
            "Mooca, São Paulo SP representado pela sócia proprietária Damaris Talita Macedo, "
            "portador do RG: 30.315.655-7."
        )

        document.add_paragraph()

        # =====================================================================
        # CLÁUSULA 1 - PRODUTOS CONTRATADOS  (BLOCO CORRIGIDO)
        # =====================================================================
        document.add_heading('CLÁUSULA 1 - PRODUTOS CONTRATADOS', level=1)

        # Tenta várias formas de vir os produtos (string JSON ou lista)
        produtos_raw = (
            dados.get('produtosContratadosJson')
            or dados.get('ProdutosContratados')
            or dados.get('Produtos Contratados')
            or []
        )

        produtos = []

        # Caso já venha como lista (mais comum no sistema)
        if isinstance(produtos_raw, list):
            produtos = produtos_raw

        # Caso venha como string JSON (mais comum na extração do contrato)
        elif isinstance(produtos_raw, str):
            try:
                produtos = json.loads(produtos_raw)
            except Exception as e:
                print(f"[AVISO] JSON de produtos inválido ao gerar DOCX: {e}")

        # Monta a tabela se tiver produtos
        if produtos:
            tabela = document.add_table(rows=1, cols=4)
            tabela.style = 'Table Grid'

            hdr_cells = tabela.rows[0].cells
            hdr_cells[0].text = 'Quantidade'
            hdr_cells[1].text = 'Produto'
            hdr_cells[2].text = 'Valor Unitário'
            hdr_cells[3].text = 'Valor Total'

            for item in produtos:
                row_cells = tabela.add_row().cells
                row_cells[0].text = str(item.get('Quantidade', ''))
                row_cells[1].text = str(item.get('Produto', ''))
                # tenta pegar tanto o nome usado no extrator quanto no sistema
                row_cells[2].text = str(item.get('Valor Unitário', item.get('valorUnitario', '')))
                row_cells[3].text = str(item.get('Valor Total Item', item.get('totalItem', '')))

            document.add_paragraph(
                f"TOTAL: {dados.get('Valor Total do Pedido', dados.get('Valor_Total_do_Pedido', 'N/A'))}"
            )
        else:
            document.add_paragraph("Nenhum produto adicionado.")

        # =====================================================================
        # RESTO DAS CLÁUSULAS (IGUAL AO SEU CÓDIGO ORIGINAL)
        # =====================================================================
        document.add_heading('CLÁUSULA 2 - VALOR E FORMA DE PAGAMENTO', level=1)
        document.add_paragraph(
            f"O valor total de {dados.get('Valor Total do Pedido', dados.get('Valor_Total_do_Pedido', 'N/A'))} "
            f"referente aos produtos acima citados, foram pagos no dia "
            f"{dados.get('Data de Pagamento', dados.get('Data_de_Pagamento', 'N/A'))} "
            f"{dados.get('Forma de Pagamento', dados.get('FormaDePagamento', 'N/A'))}."
        )

        document.add_heading('CLÁUSULA 3 - EMBALAGEM DOS DOCES - FORMINHAS', level=1)
        document.add_paragraph(
            'Os doces finos são entregues em forminhas no formato caixeta, na cor branca, todos decorados '
            'e prontos para o consumo. Os brigadeiros serão entregues em forminhas na cor branca nº 5.'
        )
        document.add_paragraph(
            'Caso o CONTRATANTE opte por embalagens decorativas, o mesmo deverá enviar ao CONTRATADO '
            'com no máximo 15 dias de antecedência ao evento, que entregará os doces finos dentro das '
            'embalagens decoradas, prontos para o consumo. Após esse prazo não recebemos.'
        )
        document.add_paragraph(
            'Por haver um manejo especial nas forminhas no modelo de flor e um custo maior de compra de '
            'caixas para armazenamento dos doces, é cobrado uma taxa adicional de R$0,10 por unidade, '
            'como consta abaixo:'
        )
        document.add_paragraph(
            'ATÉ 100 DOCES + R$10,00 / ATÉ 200 DOCES + R$20,00 ATÉ 300 DOCES + R$30,00 / '
            'ATÉ 400 DOCES + R$40,00 ACIMA DE 500 DOCES + R$50,00 e assim sucessivamente'
        )
        document.add_paragraph()

        document.add_heading('CLÁUSULA 4 - EMBALAGENS DOS BEM-CASADOS', level=1)
        document.add_paragraph(
            'Os bem-casados são entregues em papel crepom crepe plus, com celofane e fita de cetim de 7mm, '
            'nas cores enviadas na tabela completa. Os papéis perolados da linha especial serão cobrados '
            'R$ 0,40 a mais por unidade e os papéis dourado, prata, tiffany e marsala serão cobrados '
            'R$ 0,20 a mais por unidade, por se tratar de um papel especial e com maior custo. '
            'Tudo está discriminado na tabela de cores.'
        )
        document.add_paragraph(
            'Caso o CONTRATANTE opte por incluir, medalhinhas, tercinhos, renda, juta, tag ou outro item '
            'decorativo, deverá consultar antecipadamente a disponibilidade e todos os itens são colados '
            'com cola quente. A entrega dos itens deverá ocorrer com no máximo 15 dias antes do evento. '
            'Após esse prazo não recebemos. Por haver um manejo especial dos itens, será cobrado uma taxa '
            'adicional de R$0,10, como consta abaixo: ATÉ 100 BEM-CASADOS + R$10,00 / ATÉ 200 BEM-CASADOS '
            '+ R$20,00 ATÉ 300 BEM-CASADOS + R$30,00 / ATÉ 400 BEM-CASADOS + R$40,00 ACIMA de 500 '
            'BEM-CASADOS + R$50,00 e assim sucessivamente'
        )
        document.add_paragraph('Caso opte pela aplicação de dois ou mais itens, será cobrado o valor de cada item.')
        document.add_paragraph()

        document.add_heading('CLÁUSULA 5 - ALTERAÇÕES', level=1)
        document.add_paragraph(
            'Não recebemos forminhas, modificações, alterações em contrato em hipótese alguma na semana do evento.'
        )
        document.add_paragraph()

        document.add_heading('CLÁUSULA 6 - ADIÇÃO DE NOVOS ITENS', level=1)
        document.add_paragraph(
            'Caso haja a necessidade do CONTRATANTE adicionar novos itens ao pedido fechado, o valor dos '
            'produtos será de acordo com o valor vigente no momento da adição, mesmo que o contrato tenha '
            'sido fechado com valores promocionais.'
        )
        document.add_paragraph(
            'A adição de produtos ocorre de acordo com a disponibilidade de agenda. Não havendo disponibilidade '
            'para novos produtos ou pedidos, não será possível a complementação.'
        )
        document.add_paragraph()

        document.add_heading('CLÁUSULA 7 - RETIRada OU SERVIÇO DE ENTREGA', level=1)
        document.add_paragraph(
            "A entrega ou retirada dos itens acima, deverá ser definida pela CONTRATANTE até 15 dias antes do "
            "evento. Em caso de entrega será cobrada taxa de deslocamento de R$ 6,00 por km ou a taxa mínima "
            "de R$50,00 (sujeito a disponibilidade na data e horário desejados). Não fazemos entregas aos "
            "domingos e feriados. A retirada dos produtos ocorre de segunda-feira à sábado, das 9h às 16h30, "
            "mediante agendamento com o setor responsável, não havendo expediente aos domingos e feriados."
        )
        document.add_paragraph()

        document.add_heading('CLÁUSULA 8 - ARMAZENAMENTO', level=1)
        document.add_paragraph(
            'Todos os doces e/ou bem-casados deverão, obrigatoriamente, ser armazenados em geladeira até o '
            'momento da montagem da mesa para o evento. Validade 3 a 5 dias em geladeira.'
        )
        document.add_paragraph()

        document.add_heading('CLÁUSULA 9 - LOCAÇÃO (SE HOUVER)', level=1)
        document.add_paragraph(
            'Caso haja locação de bolo cenográfico, o CONTRATANTE deverá deixar uma caução no valor de R$300,00 '
            'ou o valor em dinheiro, como forma de garantia. O bolo cenográfico sendo locado e deverá retornar '
            'nas mesmas condições, em até 4 dias após a data da retirada. Na devolução do bolo cenográfico, '
            'será devolvido o valor total. Em caso de avarias será cobrado R$ 100,00 por andar (dependendo do '
            'modelo) para refazer cada andar danificado. O CONTRATANTE deverá tomar todos os cuidados necessários '
            'como: não expor ao calor excessivo,  água  ou  qualquer  outro  líquido,  não  deverá  apertar,  '
            'amassar,  não  deixar convidados colocarem as mãos e deverá ser transportado com cuidado, pegando '
            'somente pela base de madeira.'
        )
        document.add_paragraph()

        document.add_heading('CLÁUSULA 10 - REMARCAÇÃO', level=1)
        document.add_paragraph(
            'Em caso de REMARCAÇÃO de data do evento superior a 6 meses, será cobrado um reequilíbrio econômico '
            'e financeiro de 10% sobre o valor do contrato, a cada 6 meses de diferença da data marcada inicialmente.'
        )
        document.add_paragraph()

        document.add_heading('CLÁUSULA 11 - DATA E LOCAL DO EVENTO', level=1)
        document.add_paragraph(
            f"O evento acontecerá no dia: {dados.get('Data do Evento', dados.get('Data_do_Evento', 'N/A'))} "
            f"- Local do evento: {dados.get('Local do Evento', dados.get('Local_do_Evento', 'N/A'))}"
        )
        document.add_paragraph(f"Como nos conheceu: {dados.get('Como nos conheceu', 'N/A')}")
        document.add_paragraph()

        document.add_heading('CLÁUSULA 12 - CANCELAMENTO', level=1)
        document.add_paragraph(
            'A CONTRATANTE pagará multa de 30% do valor do contrato em caso de cancelamento. O CONTRATADO '
            'pagará multa de 100% do valor do contrato em caso de cancelamento.'
        )
        document.add_paragraph()

        document.add_paragraph(f"RESPONSÁVEL PELO CONTRATO: {dados.get('Responsavel', 'N/A')}")
        document.add_paragraph(
            f"São Paulo, {dados.get('Data de Pagamento', dados.get('Data_de_Pagamento', 'N/A'))}"
        )
        document.add_paragraph()

        document.add_paragraph('CONTRATANTE', style='Normal').bold = True
        document.add_paragraph('______________________________', style='Normal')
        document.add_paragraph('CONTRATADO', style='Normal').bold = True
        document.add_paragraph('______________________________', style='Normal')

        doc_stream = BytesIO()
        document.save(doc_stream)
        doc_stream.seek(0)
        return doc_stream

    except Exception as e:
        print(f"[ERRO] Falha ao gerar contrato DOCX: {e}")
        return None


def gerar_contrato_pdf_direto(dados: Dict[str, Any]) -> Optional[BytesIO]:
    try:
        html_string = render_template("contrato_template.html", dados=dados)
        pdf_bytes = HTML(string=html_string).write_pdf()
        pdf_stream = BytesIO(pdf_bytes)
        pdf_stream.seek(0)
        return pdf_stream
    except Exception as e:
        print(f"[ERRO] Falha ao gerar contrato PDF com WeasyPrint: {e}")
        return None

def gerar_relatorio_entrega(dados: Dict[str, Any]) -> Optional[BytesIO]:
    try:
        document = Document()
        document.add_heading('RELATÓRIO DE ENTREGA', 0)
        contratante = dados.get('Contratante', {})
        data_evento = dados.get('Data do Evento', dados.get('Data_do_Evento', 'Não informada'))
        local_evento = dados.get('Local do Evento', dados.get('Local_do_Evento', 'Não informado'))
        document.add_paragraph(f"Nome do Cliente: {contratante.get('Nome', 'Não encontrado')}")
        document.add_paragraph(f"Data do Evento: {data_evento}")
        document.add_paragraph(f"Local do Evento: {local_evento}")
        document.add_paragraph(f"Data de Emissão: {datetime.datetime.now().strftime('%d/%m/%Y')}")
        document.add_paragraph("\nProdutos Contratados:")
        produtos = []
        produtos_json_str = dados.get('produtosContratadosJson', '[]')
        try:
            produtos = json.loads(produtos_json_str)
        except json.JSONDecodeError:
            print("[AVISO] JSON de produtos inválido ao gerar Relatório.")
            produtos = dados.get('Produtos Contratados', [])
        if produtos:
            tabela = document.add_table(rows=1, cols=4)
            tabela.style = 'Table Grid'
            hdr_cells = tabela.rows[0].cells
            hdr_cells[0].text, hdr_cells[1].text, hdr_cells[2].text, hdr_cells[3].text = 'Quantidade', 'Produto', 'Valor Unitário', 'Valor Total'
            for item in produtos:
                row_cells = tabela.add_row().cells
                row_cells[0].text, row_cells[1].text, row_cells[2].text, row_cells[3].text = str(item.get('Quantidade', '')), str(item.get('Produto', '')), str(item.get('Valor Unitário', '')), str(item.get('Valor Total Item', ''))
        else:
            document.add_paragraph("Nenhum produto encontrado.")
        document.add_paragraph(f"\nValor Total do Pedido: R$ {dados.get('Valor Total do Pedido', dados.get('Valor_Total_do_Pedido', 'N/A'))}")
        document.add_paragraph("\n\n\nAssinaturas:\n")
        document.add_paragraph("______________________________\nResponsável pela Entrega")
        document.add_paragraph("\n\n")
        document.add_paragraph("______________________________\nResponsável pela Retirada")
        doc_stream = BytesIO()
        document.save(doc_stream)
        doc_stream.seek(0)
        return doc_stream
    except Exception as e:
        print(f"[ERRO] Falha ao salvar relatório de entrega: {e}")
        return None

def exportar_para_excel(dados: Dict[str, Any]) -> Optional[BytesIO]:
    try:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Dados do Contrato"
        sheet['A1'] = "Campo"
        sheet['B1'] = "Informação Extraída"
        linha_atual = 2
        for chave, valor in dados.items():
            if chave == 'Produtos Contratados' or chave == 'produtosContratadosJson':
                continue
            if isinstance(valor, dict):
                for sub_chave, sub_valor in valor.items():
                    sheet[f'A{linha_atual}'] = f"{chave} - {sub_chave}"
                    sheet[f'B{linha_atual}'] = sub_valor
                    linha_atual += 1
            else:
                sheet[f'A{linha_atual}'] = chave
                sheet[f'B{linha_atual}'] = valor
                linha_atual += 1
        linha_atual += 2
        produtos_contratados_str = dados.get('produtosContratadosJson')
        if produtos_contratados_str:
            produtos_contratados = json.loads(produtos_contratados_str)
            if produtos_contratados and isinstance(produtos_contratados, list) and len(produtos_contratados) > 0:
                headers_produtos = list(produtos_contratados[0].keys())
                for col_idx, header in enumerate(headers_produtos, 1):
                    sheet.cell(row=linha_atual, column=col_idx, value=header)
                linha_atual += 1
                for produto in produtos_contratados:
                    for col_idx, header in enumerate(headers_produtos, 1):
                        sheet.cell(row=linha_atual, column=col_idx, value=produto.get(header, 'N/A'))
                    linha_atual += 1
        excel_stream = BytesIO()
        workbook.save(excel_stream)
        excel_stream.seek(0)
        return excel_stream
    except Exception as e:
        print(f"\n[ERRO] Não foi possível salvar a planilha: {e}")
        return None

# ==============================================================================
# SEÇÃO 4: BLOCO DE TESTE (Inalterado)
# ==============================================================================
if __name__ == "__main__":
    caminho_do_pdf = "modelo_contrato.pdf" 
    print(f"Tentando extrair texto de: {caminho_do_pdf}")
    
    try:
        with open(caminho_do_pdf, 'rb') as f:
            pdf_bytes_content = f.read()
        
        print("\n--- TESTANDO MODO SISTEMA (REGEX) ---")
        dados_do_contrato = extrair_dados_do_contrato_por_tipo(pdf_bytes_content, tipo_analise='sistema', file_extension='.pdf') 

    except FileNotFoundError:
        print(f"\n[ERRO] Arquivo de teste '{caminho_do_pdf}' não encontrado. Crie um para testar.")
        dados_do_contrato = None
    except Exception as e:
        print(f"\n[ERRO] Falha ao ler PDF de teste: {e}")
        dados_do_contrato = None

    if dados_do_contrato:
        print("\nDados extraídos com sucesso. Imprimindo e gerando arquivos...")
        
        print("\n--- DADOS EXTRAÍDOS DO CONTRATO ---")
        for chave, valor in dados_do_contrato.items():
            print(f"\n>> {chave}:")
            if chave == 'produtosContratadosJson':
                try:
                    produtos = json.loads(valor)
                    for item in produtos:
                        print(f"    - {item}")
                except json.JSONDecodeError:
                    print(f"    {valor}")
            elif isinstance(valor, dict):
                for sub_chave, sub_valor in valor.items():
                    print(f"     {sub_chave}: {sub_valor}")
            else:
                print(f"    {valor}")
        print("\n----------------------------------------------------")

        nome_excel = f"dados_contrato_{datetime.datetime.now().strftime('%Ym%d_%H%M%S')}.xlsx"
        excel_stream = exportar_para_excel(dados_do_contrato)
        if excel_stream:
            with open(nome_excel, 'wb') as f:
                f.write(excel_stream.getvalue())
            print(f"\n[SUCESSO] Dados exportados para o arquivo '{nome_excel}'")
        else:
            print(f"\n[ERRO] Não foi possível gerar a planilha Excel.")

        nome_docx = f"relatorio_entrega_{datetime.datetime.now().strftime('%Y%H%S')}.docx"
        docx_stream = gerar_relatorio_entrega(dados_do_contrato)
        if docx_stream:
            with open(nome_docx, 'wb') as f:
                f.write(docx_stream.getvalue())
            print(f"[OK] Relatório de entrega salvo em: {nome_docx}")
        else:
            print(f"[ERRO] Falha ao salvar relatório de entrega.")

    else:
        print("\nNão foi possível processar o PDF.")
